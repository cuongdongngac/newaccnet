from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from urllib.parse import urlencode
from pathlib import Path
from typing import Any
import json
import time

from openpyxl import load_workbook

from src.gdt.client import endpoint_url, error_hint
from src.invoice.keys import invoice_key
from src.invoice.models import Window, split_calendar_months
from src.invoice.parser import extract_count, extract_rows


class InvoiceScanner:
    """Production scanner: query -> export cross-check -> detail -> artifact detection."""

    ACTIONS = {
        ('purchase', 'hddt'): 'Xem hóa đơn (hóa đơn điện tử mua vào)',
        ('sold', 'hddt'): 'Xem hóa đơn (hóa đơn điện tử bán ra)',
        ('purchase', 'pos'): 'Xem hóa đơn (hóa đơn máy tính tiền mua vào)',
        ('sold', 'pos'): 'Xem hóa đơn (hóa đơn máy tính tiền bán ra)',
    }

    def __init__(self, cfg, client, storage, logger, owner_mst: str = '', progress=None):
        self.cfg = cfg
        self.owner_mst = str(owner_mst or '').strip()
        self.client = client
        self.storage = storage
        self.log = logger
        self.progress = progress or (lambda message: None)
        self.manifest = {
            'scanner_version': '3.1.2',
            'owner_mst': self.owner_mst,
            'invoices': {},
            'windows': [],
            'summary': {},
        }

    def query_cfg(self):
        return self.cfg['invoice_query']

    def build_search(self, window: Window) -> str:
        q = self.query_cfg()
        fmt = q['date']['format']
        a = datetime.combine(window.start, datetime.min.time()).strftime(fmt)
        b = datetime.combine(window.end, datetime.max.time().replace(microsecond=0)).strftime(fmt)
        return q.get('search_template', 'tdlap=ge={from_date};tdlap=le={to_date}').format(from_date=a, to_date=b).strip(';')

    def categories(self):
        out = []
        for side in ('purchase', 'sold'):
            for category in ('hddt', 'pos'):
                item = self.cfg.get('invoice', {}).get(side, {}).get(category, {})
                if item.get('enabled', False):
                    out.append((side, category, item))
        return out

    def _query_url(self, item, side, window, page):
        q = self.query_cfg()
        url = endpoint_url(self.cfg, item['endpoint'])
        sort_value = q.get('sort', {}).get(side, 'tdlap:desc')
        params = [
            ('sort', sort_value),
            ('size', str(q.get('size', 15))),
            ('search', self.build_search(window)),
        ]
        if q.get('include_page', True):
            start = int(q.get('page_start', 0))
            if page > start or q.get('send_page_zero', False):
                params.append(('page', str(page)))
        # Let requests/urllib encode the search expression correctly. Raw query
        # strings are fragile when symbols such as ';', ':', spaces or unicode
        # appear and can produce different server-side parsing.
        return url + '?' + urlencode(params)

    def _query_page(self, side, category, item, window, page, deadline=None):
        if deadline is not None and time.monotonic() >= deadline:
            return {'status': 'WINDOW_TIMEOUT', 'side': side, 'category': category, 'page': page, 'diagnostic': 'Đã đạt giới hạn thời gian của cửa sổ'}
        url = self._query_url(item, side, window, page)
        headers = {'End-Point': '/tra-cuu/tra-cuu-hoa-don'}
        try:
            r = self.client.request(
                item.get('method', 'GET').upper(),
                url,
                authenticated=True,
                headers=headers,
                timeout=min(float(self.cfg['connection'].get('invoice_timeout_seconds', 25)), max(1.0, deadline-time.monotonic())) if deadline is not None else float(self.cfg['connection'].get('invoice_timeout_seconds', 25)),
            )
        except Exception as exc:
            return {
                'status': 'network_error', 'diagnostic': str(exc)[:500],
                'side': side, 'category': category, 'page': page,
                'window': {'from': window.start.isoformat(), 'to': window.end.isoformat()},
            }

        result = {
            'status': 'ok' if r.status_code == 200 else 'http_error',
            'http_status': r.status_code,
            'side': side,
            'category': category,
            'page': page,
            'window': {'from': window.start.isoformat(), 'to': window.end.isoformat()},
            'query': {
                'search': self.build_search(window),
                'size': int(self.query_cfg().get('size', 15)),
                'sort': self.query_cfg().get('sort', {}).get(side, 'tdlap:desc'),
                'page': page,
            },
        }
        if r.status_code != 200:
            result['diagnostic'] = error_hint(r)
            return result
        try:
            data = r.json()
        except ValueError:
            result.update(status='invalid_json', diagnostic='HTTP 200 nhưng response không phải JSON')
            return result

        rows = extract_rows(data)
        result['response_shape'] = 'list' if isinstance(data, list) else sorted(map(str, data.keys()))[:40] if isinstance(data, dict) else type(data).__name__
        result['count_hint'] = extract_count(data)
        result['page_row_count'] = len(rows) if rows is not None else None
        if rows is None:
            result['diagnostic'] = 'Không tìm thấy mảng bản ghi hóa đơn trong JSON response'
            return result
        result['invoice_file'] = str(self.storage.save_page(side, category, window, page, rows))
        result['invoice_rows_saved'] = len(rows)
        result['xml_files_saved'] = self.storage.save_embedded_xml(side, category, window, rows)
        return result

    def discover(self, side, category, item, window, deadline=None):
        size = int(self.query_cfg().get('size', 15))
        max_pages = int(self.query_cfg().get('max_pages', 100))
        start = int(self.query_cfg().get('page_start', 0))
        retry_limit = int(self.query_cfg().get('page_retries', 2))
        pages = []
        rows_by_key: dict[str, dict[str, Any]] = {}
        expected = None
        repeated = False
        page = start

        for _ in range(max_pages):
            res = None
            for attempt in range(1, retry_limit + 1):
                if deadline is not None and time.monotonic() >= deadline:
                    res = {'status':'WINDOW_TIMEOUT','page':page,'diagnostic':'Đã đạt giới hạn thời gian cửa sổ'}
                    break
                res = self._query_page(side, category, item, window, page, deadline=deadline)
                if res.get('status') == 'ok':
                    break
                status = res.get('http_status')
                if status not in (403, 408, 429, 500, 502, 503, 504) and not str(res.get('status', '')).endswith('error'):
                    break
                if attempt < retry_limit:
                    wait = float(self.query_cfg().get('throttle_retry_seconds', 1.5)) * attempt
                    if deadline is not None: wait = min(wait, max(0.0, deadline-time.monotonic()))
                    time.sleep(wait)
            assert res is not None
            pages.append(res)
            if res.get('status') != 'ok':
                break

            try:
                rows = json.loads(Path(res['invoice_file']).read_text(encoding='utf-8')).get('invoices', [])
            except Exception:
                rows = []
            before = len(rows_by_key)
            for row in rows:
                if isinstance(row, dict):
                    rows_by_key.setdefault(invoice_key(row, side, category), row)
            new_count = len(rows_by_key) - before
            if rows and new_count == 0:
                repeated = True
                break

            count = res.get('count_hint')
            if isinstance(count, int) and count >= 0:
                expected = count if expected is None else max(expected, count)
            if not rows or len(rows) < size:
                break
            if expected is not None and len(rows_by_key) >= expected:
                break
            page += 1

        return {'rows': list(rows_by_key.values()), 'expected': expected, 'pages': pages, 'repeated': repeated}

    def export_verify(self, side, category, item, window, deadline=None):
        ecfg = item.get('export') or {}
        if not ecfg.get('enabled', False):
            return {'enabled': False, 'rows': [], 'status': 'disabled', 'count': 0}
        url = endpoint_url(self.cfg, ecfg['endpoint'])
        search = self.build_search(window)
        params = urlencode([('sort', ecfg.get('sort', 'tdlap:desc')), ('search', search)])
        if ecfg.get('extra_query'):
            extra = str(ecfg['extra_query']).lstrip('&?')
            params += '&' + extra
        url += '?' + params
        headers = {
            'Action': ecfg.get('action', 'Xuất hóa đơn'),
            'End-Point': '/tra-cuu/tra-cuu-hoa-don',
            'Accept': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream',
        }
        last = ''
        retries = int(ecfg.get('retries', 1))
        for attempt in range(1, retries + 1):
            try:
                r = self.client.request(
                    'GET', url, authenticated=True, headers=headers,
                    timeout=min(float(self.cfg['connection'].get('invoice_timeout_seconds', 25)), max(1.0, deadline-time.monotonic())) if deadline is not None else float(self.cfg['connection'].get('invoice_timeout_seconds', 25)),
                )
            except Exception as exc:
                last = str(exc)
                if attempt < retries:
                    time.sleep(attempt)
                continue
            if r.status_code == 200:
                try:
                    wb = load_workbook(BytesIO(r.content), read_only=True, data_only=True)
                    if not wb.sheetnames:
                        wb.close()
                        return {'enabled': True, 'status': 'empty_workbook', 'rows': [], 'count': 0}
                    ws = wb[wb.sheetnames[0]]
                    data = list(ws.iter_rows(values_only=True))
                    wb.close()
                    header = next((i for i, row in enumerate(data) if row and str(row[0]).strip() == 'STT'), None)
                    if header is None:
                        return {'enabled': True, 'status': 'no_header', 'rows': [], 'count': 0}
                    cm = ecfg.get('column_map', {})
                    out = []
                    for row in data[header + 1:]:
                        if not any(v not in (None, '') for v in row):
                            continue
                        out.append({k: (row[idx] if idx < len(row) else '') for k, idx in cm.items()})
                    return {'enabled': True, 'status': 'ok', 'rows': out, 'count': len(out)}
                except Exception as exc:
                    last = f'Excel parse: {exc}'
            else:
                last = f'HTTP {r.status_code}: {error_hint(r)}'
            if attempt < retries:
                time.sleep(max(1.0, attempt * 1.5 if r.status_code not in (403, 429) else attempt * 2.0))
        return {'enabled': True, 'status': 'failed', 'rows': [], 'count': 0, 'error': last[:700]}

    @staticmethod
    def _detail_is_error(data: Any) -> str | None:
        if not isinstance(data, dict):
            return None
        # GDT sometimes returns HTTP 200 with an application-level error object.
        message = data.get('message') or data.get('error') or data.get('errorMessage')
        if isinstance(message, str) and message.strip():
            return message.strip()[:500]
        return None

    def fetch_detail(self, side, category, item, window, row, deadline=None):
        dcfg = item.get('detail') or {}
        if not dcfg.get('enabled', True):
            return None, 'DISABLED', ''
        params = {
            'nbmst': str(row.get('nbmst') or ''),
            'khhdon': str(row.get('khhdon') or ''),
            'shdon': str(row.get('shdon') or ''),
            'khmshdon': str(row.get('khmshdon') or ''),
        }
        if not all(params.values()):
            return None, 'SKIPPED_MISSING_KEY', 'Thiếu nbmst/khhdon/shdon/khmshdon'

        url = endpoint_url(self.cfg, dcfg.get('endpoint', '/api/query/invoices/detail'))
        delay = float(dcfg.get('delay_seconds', 0.35))
        retries = int(dcfg.get('retries', 2))
        last = ''
        for attempt in range(1, retries + 1):
            if deadline is not None and time.monotonic() >= deadline:
                return None, 'WINDOW_TIMEOUT', 'Đã đạt giới hạn thời gian cửa sổ'
            # Pace requests, but never sleep past the window deadline.
            wait = delay if attempt == 1 else float(dcfg.get('backoff_seconds', 0.75)) * attempt
            if deadline is not None: wait = min(wait, max(0.0, deadline-time.monotonic()))
            time.sleep(wait)
            try:
                r = self.client.request(
                    'GET', url,
                    authenticated=True,
                    headers={
                        'Action': dcfg.get('action', self.ACTIONS[(side, category)]),
                        'End-Point': '/tra-cuu/tra-cuu-hoa-don',
                    },
                    params=params,
                    timeout=min(float(self.cfg['connection'].get('invoice_timeout_seconds', 25)), max(1.0, deadline-time.monotonic())) if deadline is not None else float(self.cfg['connection'].get('invoice_timeout_seconds', 25)),
                )
            except Exception as exc:
                last = str(exc)
                continue

            if r.status_code == 200:
                try:
                    data = r.json()
                except ValueError:
                    last = 'HTTP 200 nhưng detail không phải JSON'
                    continue
                app_error = self._detail_is_error(data)
                if app_error:
                    last = f'application_error: {app_error}'
                    continue
                if data in (None, {}, []):
                    last = 'HTTP 200 nhưng detail rỗng'
                    continue
                return data, 'OK', ''

            last = f'HTTP {r.status_code}: {error_hint(r)}'
            if r.status_code == 401:
                return None, 'AUTH_ERROR', last
            if r.status_code not in (403, 408, 429, 500, 502, 503, 504):
                break

        return None, 'FAILED', last[:700]

    def _persist_progress(self, start: date, end: date, summary: dict, results: list):
        self.manifest['scan_range'] = {'from': start.isoformat(), 'to': end.isoformat()}
        self.manifest['summary'] = summary
        self.manifest['results_count'] = len(results)
        self.manifest['status'] = 'RUNNING'
        self.storage.write_manifest(self.manifest)

    def run(self, start: date, end: date):
        q = self.query_cfg()
        windows = split_calendar_months(start, end)
        all_results = []
        summary = {}
        for side, category, item in self.categories():
            key = f'{side}_{category}'
            side_stats = {
                'windows': 0, 'pages': 0, 'discovered': 0, 'expected': 0,
                'export_count': 0, 'missing_from_query': 0,
                'details_ok': 0, 'details_failed': 0, 'xml_ok': 0,
                'query_errors': 0, 'export_errors': 0, 'pagination_stalled': 0,
                'incomplete_windows': 0, 'retries_recovered': 0,
            }
            for window in windows:
                side_stats['windows'] += 1
                window_timeout = float(self.query_cfg().get('window_timeout_seconds', 90))
                phase_deadline = time.monotonic() + window_timeout
                self.progress(f'{side.upper()} / {category.upper()} — {window.start:%d/%m/%Y} → {window.end:%d/%m/%Y} — bắt đầu')
                d = self.discover(side, category, item, window, deadline=phase_deadline)
                side_stats['pages'] += len(d['pages'])
                side_stats['discovered'] += len(d['rows'])
                if any(p.get('status') != 'ok' for p in d['pages']):
                    side_stats['query_errors'] += 1
                if d.get('repeated'):
                    side_stats['pagination_stalled'] += 1

                export_deadline = time.monotonic() + float(self.query_cfg().get('export_timeout_seconds', 25))
                if time.monotonic() >= phase_deadline:
                    exp = {'enabled': bool((item.get('export') or {}).get('enabled', False)), 'rows': [], 'status': 'QUERY_PHASE_TIMEOUT', 'count': 0}
                else:
                    self.progress(f'{side.upper()} / {category.upper()} — {window.start:%m/%Y} — đối soát Export')
                    exp = self.export_verify(side, category, item, window, deadline=export_deadline)
                if exp.get('status') not in ('ok', 'disabled'):
                    side_stats['export_errors'] += 1
                side_stats['export_count'] += int(exp.get('count') or 0)

                export_keys = {invoice_key(r, side, category) for r in exp.get('rows', []) if isinstance(r, dict)}
                query_keys = {invoice_key(r, side, category) for r in d['rows'] if isinstance(r, dict)}
                missing = export_keys - query_keys
                side_stats['missing_from_query'] += len(missing)

                merged = {invoice_key(r, side, category): r for r in d['rows'] if isinstance(r, dict)}
                for r in exp.get('rows', []):
                    if isinstance(r, dict):
                        merged.setdefault(invoice_key(r, side, category), r)
                query_count = len(query_keys)
                export_count = len(export_keys) if exp.get('status') == 'ok' else 0
                count_hints = [int(d.get('expected') or 0), query_count, export_count]
                expected = max(count_hints)
                side_stats['expected'] += expected

                # A month is considered complete only when the independent
                # discovery sources agree. If export is unavailable, the list
                # count is still usable but is explicitly marked as a weaker
                # reconciliation state. Never silently call a partial month OK.
                recon = 'OK'
                if any(p.get('status') != 'ok' for p in d['pages']):
                    recon = 'QUERY_ERROR'
                elif d.get('repeated'):
                    recon = 'PAGINATION_STALLED'
                elif exp.get('status') == 'ok' and query_count != export_count:
                    recon = 'COUNT_MISMATCH'
                elif d.get('expected') not in (None, query_count):
                    recon = 'COUNT_MISMATCH'
                elif exp.get('status') not in ('ok', 'disabled'):
                    recon = 'EXPORT_UNAVAILABLE'
                if recon != 'OK':
                    side_stats['incomplete_windows'] += 1

                # IMPORTANT: do not reuse the discovery/export deadline for Detail.
                # A slow Export must never consume the Detail budget and turn every
                # discovered invoice into WINDOW_TIMEOUT. Detail gets its own bounded
                # phase budget after discovery is complete.
                detail_deadline = time.monotonic() + float(self.query_cfg().get('detail_phase_timeout_seconds', 120))
                failed_rows = []
                detail_total = len(merged)
                detail_done = 0
                for row in merged.values():
                    detail_done += 1
                    self.progress(f'{side.upper()} / {category.upper()} — {window.start:%m/%Y} — Detail {detail_done}/{detail_total}')
                    p = self.storage.upsert_invoice(side, category, window, row, source='query', detail_status='PENDING', error=None)
                    detail, st, err = self.fetch_detail(side, category, item, window, row, deadline=detail_deadline)
                    source = 'export' if invoice_key(row, side, category) in missing else 'query'
                    p = self.storage.upsert_invoice(
                        side, category, window, row,
                        detail=detail,
                        source=source,
                        detail_status=st,
                        error='' if st == 'OK' else err,
                    )
                    if st != 'OK':
                        failed_rows.append(row)
                    try:
                        record = json.loads(p.read_text(encoding='utf-8'))
                    except Exception:
                        record = {}
                    rec = {
                        'key': invoice_key(row, side, category),
                        'side': side,
                        'category': category,
                        'detail_status': st,
                        'xml_status': record.get('xml_status', 'NOT_FOUND'),
                        'detail_file': record.get('detail_file', ''),
                        'xml_file': record.get('xml_file', ''),
                        'reconciliation': recon,
                        'error': err,
                        'record_file': str(p.relative_to(self.storage.root)),
                    }
                    if st == 'OK':
                        side_stats['details_ok'] += 1
                    else:
                        side_stats['details_failed'] += 1
                    all_results.append(rec)

                # One deliberately conservative recovery pass. The user values
                # completeness over speed; a transient GDT failure must not
                # permanently become a missing detail when a second pass can
                # recover it. Successful records are never downgraded.
                recovery_retries = int((item.get('detail') or {}).get('recovery_retries', 1))
                for row in failed_rows:
                    for _ in range(recovery_retries):
                        time.sleep(float((item.get('detail') or {}).get('recovery_delay_seconds', 2.0)))
                        if time.monotonic() >= detail_deadline:
                            break
                        detail, st, err = self.fetch_detail(side, category, item, window, row, deadline=detail_deadline)
                        if st != 'OK':
                            continue
                        p = self.storage.upsert_invoice(
                            side, category, window, row, detail=detail,
                            source='recovery', detail_status='OK', error=''
                        )
                        side_stats['details_failed'] = max(0, side_stats['details_failed'] - 1)
                        side_stats['details_ok'] += 1
                        side_stats['retries_recovered'] += 1
                        try:
                            record = json.loads(p.read_text(encoding='utf-8'))
                        except Exception:
                            record = {}
                        all_results.append({
                            'key': invoice_key(row, side, category), 'side': side, 'category': category,
                            'detail_status': 'OK', 'xml_status': record.get('xml_status', 'NOT_FOUND'),
                            'detail_file': record.get('detail_file', ''), 'xml_file': record.get('xml_file', ''),
                            'reconciliation': recon, 'error': '', 'record_file': str(p.relative_to(self.storage.root)),
                        })
                        break
                all_results.extend(d['pages'])
                self.manifest['windows'].append({
                    'side': side, 'category': category,
                    'from': window.start.isoformat(), 'to': window.end.isoformat(),
                    'reconciliation': recon, 'query_count': query_count,
                    'export_count': export_count, 'expected': expected,
                    'details_ok': sum(1 for r in all_results if r.get('side') == side and r.get('category') == category and r.get('detail_status') == 'OK'),
                    'details_failed': sum(1 for r in all_results if r.get('side') == side and r.get('category') == category and r.get('detail_status') != 'OK' and r.get('key')),
                })
                self._persist_progress(start, end, summary, all_results)

            summary[key] = side_stats
            self._persist_progress(start, end, summary, all_results)

        metric_keys = (
            'windows', 'pages', 'discovered', 'expected', 'export_count',
            'missing_from_query', 'details_ok', 'details_failed', 'xml_ok',
            'query_errors', 'export_errors', 'pagination_stalled', 'incomplete_windows',
            'retries_recovered',
        )
        summary['all'] = {
            k: sum(v.get(k, 0) for v in summary.values() if isinstance(v, dict))
            for k in metric_keys
        }
        self.manifest['summary'] = summary
        self.manifest['results_count'] = len(all_results)
        self.manifest['status'] = 'COMPLETED'
        self.manifest['completed_at'] = datetime.now().isoformat(timespec='seconds')
        self.storage.write_manifest(self.manifest)
        return {'windows': len(windows), 'results': all_results, 'summary': summary, 'aborted': None}
