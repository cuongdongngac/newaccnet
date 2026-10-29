from __future__ import annotations

import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

ALIASES = {
    'Ngày lập hóa đơn': ('tdlap', 'ngayLap', 'ngay_lap', 'ngayLapHoaDon', 'invoiceDate'),
    'Mã số thuế người bán': ('nbmst', 'mstNguoiBan', 'mstNB', 'sellerTaxCode'),
    'Tên người bán': ('nbten', 'nbtt', 'tenNguoiBan', 'sellerName'),
    'Mã số thuế người mua': ('nmmst', 'mstNguoiMua', 'mstNM', 'buyerTaxCode'),
    'Tên người mua': ('nmten', 'nmtnm', 'tenNguoiMua', 'buyerName'),
    'Ký hiệu mẫu số': ('khmshdon', 'khmshd', 'kyHieuMauSo'),
    'Ký hiệu hóa đơn': ('khhdon', 'khhd', 'kyHieuHoaDon'),
    'Số hóa đơn': ('shdon', 'sohdon', 'soHoaDon', 'soHoaDonHDon'),
    'Giá trước VAT': ('tgtcthue', 'tienTruocThue', 'subtotal', 'amountBeforeTax'),
    'VAT': ('tgtthue', 'tienThue', 'vatAmount', 'taxAmount'),
    'Tổng cộng': ('tgtttbso', 'tgtttoan', 'tongTienThanhToan', 'totalAmount', 'amountAfterTax'),
    'Trạng thái hóa đơn': ('tthai', 'trangThai', 'invoiceStatus'),
    'Kết quả kiểm tra': ('ttxly', 'ketQuaKiemTra', 'inspectionResult'),
    'Mã tra cứu': ('maTraCuu', 'key', 'keyHdon', 'id', 'idHDon'),
}

MONEY_HEADERS = {'Giá trước VAT', 'VAT', 'Tổng cộng'}
INTEGER_FORMAT = '#,##0;[Red]-#,##0;0'


def first(row: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    if not isinstance(row, dict):
        return ''
    for k in aliases:
        if k in row and row[k] not in (None, ''):
            return row[k]
    lower = {str(k).lower(): v for k, v in row.items()}
    for k in aliases:
        value = lower.get(k.lower(), '')
        if value not in (None, ''):
            return value
    return ''


def money(v: Any) -> int:
    """Parse GDT/Excel money robustly and return whole VND units.

    GDT exports may provide a real number, a Vietnamese-formatted string
    (1.234.567,89), or an English-formatted string (1,234,567.89).
    The report intentionally works in integer VND and never writes a
    decimal money format.
    """
    if v in (None, ''):
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, Decimal):
        return int(v.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    if isinstance(v, (int, float)):
        return int(Decimal(str(v)).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    s = str(v).strip().replace('\u00a0', '').replace(' ', '')
    if not s:
        return 0
    # Parenthesized negatives: (1.234) -> -1234
    negative = s.startswith('(') and s.endswith(')')
    if negative:
        s = s[1:-1]
    s = s.replace('%', '')
    try:
        if ',' in s and '.' in s:
            # Last separator is the decimal separator.
            if s.rfind(',') > s.rfind('.'):
                normalized = s.replace('.', '').replace(',', '.')
            else:
                normalized = s.replace(',', '')
        elif ',' in s:
            tail = s.rsplit(',', 1)[1]
            normalized = s.replace(',', '.') if len(tail) <= 2 else s.replace(',', '')
        elif '.' in s:
            tail = s.rsplit('.', 1)[1]
            normalized = s if len(tail) <= 2 else s.replace('.', '')
        else:
            normalized = s
        value = Decimal(normalized)
        value = value.quantize(Decimal('1'), rounding=ROUND_HALF_UP)
        result = int(value)
        return -result if negative else result
    except (InvalidOperation, ValueError):
        return 0


def parse_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or '').strip()
    for fmt in (
        '%d/%m/%Y', '%d/%m/%YT%H:%M:%S', '%d/%m/%Y %H:%M:%S',
        '%Y-%m-%d', '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S',
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def quarter(start: date, end: date) -> str:
    if start.year == end.year:
        if (start.month, end.month) in ((1, 3), (4, 6), (7, 9), (10, 12)):
            return f'Q{((start.month - 1) // 3) + 1}'
        if start == date(start.year, 1, 1) and end == date(start.year, 12, 31):
            return 'NAM'
    return f'{start:%Y%m%d}_{end:%Y%m%d}'


def _detail_values(item: dict[str, Any]) -> dict[str, Any]:
    detail = item.get('detail') if isinstance(item.get('detail'), dict) else {}
    base = item.get('list') if isinstance(item.get('list'), dict) else {}
    merged = dict(base)
    if detail:
        merged.update({k: v for k, v in detail.items() if v not in (None, '')})
    return merged


def _stable_key(row: dict[str, Any], side: str, category: str) -> tuple[str, ...]:
    return (
        side,
        category,
        str(row.get('nbmst') or ''),
        str(row.get('khhdon') or ''),
        str(row.get('shdon') or ''),
        str(row.get('khmshdon') or ''),
    )


def collect(invoice_root: Path, start: date, end: date, owner_mst: str = ''):
    purchase: list[dict[str, Any]] = []
    sold: list[dict[str, Any]] = []
    seen: set[tuple[str, ...]] = set()
    for fp in sorted(invoice_root.glob('**/records/*.json')):
        try:
            payload = json.loads(fp.read_text(encoding='utf-8'))
        except Exception:
            continue
        wf = parse_date(payload.get('window', {}).get('from'))
        wt = parse_date(payload.get('window', {}).get('to'))
        if wf and wt and (wt < start or wf > end):
            continue
        side = payload.get('side')
        if side not in ('purchase', 'sold'):
            continue
        row = _detail_values(payload)
        # Isolate data by the currently selected taxpayer. This is essential
        # because the user can change MST in the dashboard while the shared
        # invoices directory still contains historical records.
        if owner_mst:
            owner = str(owner_mst).strip()
            seller_mst = str(first(row, ALIASES['Mã số thuế người bán']) or '').strip()
            buyer_mst = str(first(row, ALIASES['Mã số thuế người mua']) or '').strip()
            side = payload.get('side')
            expected_owner = seller_mst if side == 'sold' else buyer_mst
            if expected_owner and expected_owner != owner:
                continue
        item = {h: first(row, a) for h, a in ALIASES.items()}
        item['Loại'] = 'Mua vào' if side == 'purchase' else 'Bán ra'
        item['Nguồn'] = 'Hóa đơn điện tử' if payload.get('category') == 'hddt' else 'Hóa đơn máy tính tiền'
        item['Trạng thái tải'] = payload.get('detail_status', 'UNKNOWN')
        item['Lỗi tải'] = payload.get('error', '')
        item['Khoảng tra cứu'] = f'{wf.isoformat() if wf else ""} → {wt.isoformat() if wt else ""}'
        for h in MONEY_HEADERS:
            item[h] = money(item[h])
        # The report's accounting total is ALWAYS derived, never trusted from
        # a possibly stale/locale-formatted source field. This makes the
        # workbook internally auditable: Tổng cộng = trước VAT + VAT.
        source_total = item['Tổng cộng']
        item['Tổng cộng'] = item['Giá trước VAT'] + item['VAT']
        item['Đối soát tổng'] = 'KHỚP' if source_total == item['Tổng cộng'] else f'NGUỒN={source_total:,}'.replace(',', '.')
        item['File XML'] = payload.get('xml_file', '')
        item['File chi tiết'] = str(fp.relative_to(invoice_root))
        key = _stable_key(row, side, str(payload.get('category') or 'hddt'))
        if key in seen:
            continue
        seen.add(key)
        (purchase if side == 'purchase' else sold).append(item)
    return purchase, sold


def _apply_integer_formats(wb) -> None:
    """Final workbook-wide pass: eliminate decimal money/number displays.

    This is deliberately done after all sheets/tables/formulas are created so
    no later style operation can reintroduce Excel's default 0.00/2-decimal
    display. Date cells are excluded.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.__class__.__name__ == 'MergedCell':
                    continue
                # Formula cells in this workbook are all numeric totals.
                if cell.data_type == 'f':
                    cell.number_format = INTEGER_FORMAT
                    continue
                if cell.is_date:
                    continue
                if isinstance(cell.value, (int, float, Decimal)) and not isinstance(cell.value, bool):
                    cell.number_format = INTEGER_FORMAT
                elif isinstance(cell.value, str) and cell.value.strip() in {'0.00', '0,00', '-0.00', '-0,00'}:
                    # Remove literal decimal text left by an upstream export.
                    cell.value = 0
                    cell.number_format = INTEGER_FORMAT


def build_excel(company: dict, invoice_root: Path, start: date, end: date) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    purchase, sold = collect(invoice_root, start, end, owner_mst=str(company.get('mst') or ''))
    mst = str(company.get('mst', '')).strip()
    name = str(company.get('company_name') or company.get('name') or '').strip()
    q = quarter(start, end)
    out = invoice_root / f'{re.sub(r"[^A-Za-z0-9._-]+", "_", mst or "MST")}_{q}_{end.year}.xlsx'

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tổng hợp'
    mua = wb.create_sheet('Mua')
    ban = wb.create_sheet('Bán')

    blue = '1F4E78'
    white = 'FFFFFF'
    border_color = 'B7C9D6'
    pos = '0000FF'
    neg = 'FF0000'
    thin = Side(style='thin', color=border_color)

    def style(cell, **kw):
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        cell.alignment = Alignment(**kw)

    ws.merge_cells('A1:D1')
    ws['A1'] = str(company.get('software_name') or 'GDT Invoice Manager')
    ws['A1'].font = Font(bold=True, size=18, color=white)
    ws['A1'].fill = PatternFill('solid', fgColor=blue)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 32
    ws.merge_cells('A2:D2')
    ws['A2'] = f'Tên công ty: {name or "Chưa khai báo"}'
    ws['A2'].font = Font(bold=True, size=12)
    ws['A3'], ws['B3'], ws['C3'], ws['D3'] = 'Mã số thuế', mst, 'Kỳ báo cáo', f'{q}/{end.year}'
    ws['A4'], ws['B4'], ws['C4'], ws['D4'] = 'Thời gian', start.strftime('%d/%m/%Y'), 'đến', end.strftime('%d/%m/%Y')
    for x in ('A3', 'C3', 'A4', 'C4'):
        ws[x].font = Font(bold=True)

    for c, h in enumerate(('Chi tiết', 'MUA VÀO', 'BÁN RA', 'Tổng hợp'), 1):
        x = ws.cell(6, c, h)
        x.font = Font(bold=True, color=white)
        x.fill = PatternFill('solid', fgColor=blue)
        x.alignment = Alignment(horizontal='center')
        style(x)

    rows = [
        ('Số hóa đơn', len(purchase), len(sold), '=B7+C7'),
        ('Giá trước VAT', sum(x['Giá trước VAT'] for x in purchase), sum(x['Giá trước VAT'] for x in sold), '=B8-C8'),
        ('VAT', sum(x['VAT'] for x in purchase), sum(x['VAT'] for x in sold), '=B9-C9'),
        # Accounting invariant required by the application: total is derived.
        ('Tổng cộng', '=B8+B9', '=C8+C9', '=D8+D9'),
    ]
    for r, (label, b, c, d) in enumerate(rows, 7):
        for col, val in enumerate((label, b, c, d), 1):
            x = ws.cell(r, col, val)
            style(x, horizontal='left' if col == 1 else 'right')
            if col > 1:
                x.number_format = INTEGER_FORMAT
        if r in (8, 9, 10):
            ws.cell(r, 4).font = Font(bold=True)

    ws.conditional_formatting.add('D8:D10', CellIsRule(operator='greaterThan', formula=['0'], font=Font(color=pos, bold=True)))
    ws.conditional_formatting.add('D8:D10', CellIsRule(operator='lessThan', formula=['0'], font=Font(color=neg, bold=True)))
    # Reconciliation block: keep the workbook self-auditing without adding a
    # fourth worksheet. These figures are calculated from the exact rows that
    # feed the report, so the dashboard and Excel share one accounting source.
    audit_row = 13
    ws.merge_cells(start_row=audit_row, start_column=1, end_row=audit_row, end_column=4)
    ws.cell(audit_row, 1, 'ĐỐI SOÁT DỮ LIỆU')
    ws.cell(audit_row, 1).font = Font(bold=True, color=white)
    ws.cell(audit_row, 1).fill = PatternFill('solid', fgColor=blue)
    for rr, (label, b, c, d) in enumerate([
        ('Số dòng chi tiết', len(purchase), len(sold), '=B14+C14'),
        ('Số dòng lệch tổng nguồn', sum(1 for x in purchase if x.get('Đối soát tổng') != 'KHỚP'), sum(1 for x in sold if x.get('Đối soát tổng') != 'KHỚP'), '=B15+C15'),
        ('Thiếu detail', sum(1 for x in purchase if x.get('Trạng thái tải') != 'OK'), sum(1 for x in sold if x.get('Trạng thái tải') != 'OK'), '=B16+C16'),
        ('Có XML', sum(1 for x in purchase if x.get('File XML')), sum(1 for x in sold if x.get('File XML')), '=B17+C17'),
    ], audit_row + 1):
        for col, val in enumerate((label, b, c, d), 1):
            cell = ws.cell(rr, col, val)
            style(cell, horizontal='left' if col == 1 else 'right')
            if col > 1:
                cell.number_format = INTEGER_FORMAT

    ws.freeze_panes = 'A7'
    for c, w in {'A': 30, 'B': 22, 'C': 22, 'D': 22}.items():
        ws.column_dimensions[c].width = w

    headers = [
        'Ngày lập hóa đơn', 'Loại', 'Nguồn', 'Mã số thuế người bán', 'Tên người bán',
        'Mã số thuế người mua', 'Tên người mua', 'Ký hiệu mẫu số', 'Ký hiệu hóa đơn',
        'Số hóa đơn', 'Giá trước VAT', 'VAT', 'Tổng cộng', 'Trạng thái hóa đơn',
        'Kết quả kiểm tra', 'Mã tra cứu', 'Trạng thái tải', 'Lỗi tải', 'Khoảng tra cứu',
        'Đối soát tổng', 'File XML', 'File chi tiết',
    ]

    def detail(sheet, title, data_rows, table_name):
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.cell(1, 1, title)
        sheet.cell(1, 1).font = Font(bold=True, size=15, color=white)
        sheet.cell(1, 1).fill = PatternFill('solid', fgColor=blue)
        sheet.cell(1, 1).alignment = Alignment(horizontal='center')
        for col, h in enumerate(headers, 1):
            x = sheet.cell(3, col, h)
            x.font = Font(bold=True, color=white)
            x.fill = PatternFill('solid', fgColor=blue)
            x.alignment = Alignment(horizontal='center', wrap_text=True)
            style(x)
        for rr, item in enumerate(data_rows, 4):
            for cc, h in enumerate(headers, 1):
                value = item.get(h, '')
                x = sheet.cell(rr, cc, value)
                style(x, vertical='top', wrap_text=h in ('Lỗi tải', 'Tên người bán', 'Tên người mua'))
                if h in MONEY_HEADERS:
                    x.number_format = INTEGER_FORMAT
        sheet.freeze_panes = 'A4'
        sheet.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{max(3, len(data_rows) + 3)}'
        widths = [18, 12, 24, 18, 30, 18, 30, 16, 16, 14, 18, 18, 18, 22, 24, 24, 18, 40, 24, 16, 40, 45]
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        if data_rows:
            tab = Table(displayName=table_name, ref=f'A3:{get_column_letter(len(headers))}{len(data_rows) + 3}')
            tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(tab)

    detail(mua, 'HÓA ĐƠN MUA VÀO', purchase, 'MuaHoaDon')
    detail(ban, 'HÓA ĐƠN BÁN RA', sold, 'BanHoaDon')

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
    _apply_integer_formats(wb)
    wb.save(out)
    return out
