from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import json, sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

from src.invoice.models import split_windows, split_calendar_months
from src.invoice.parser import extract_rows,extract_count,find_xml
from src.invoice.keys import invoice_key
from src.invoice.storage import InvoiceStorage
from src.invoice.scanner import InvoiceScanner
from src.reports.excel import quarter,collect


def test_split():
    w=split_windows(date(2026,7,1),date(2026,9,30),31)
    assert len(w)==3 and w[0].start==date(2026,7,1) and w[-1].end==date(2026,9,30)


def test_parser():
    d={'data':{'content':[{'id':1}]},'totalElements':1}
    assert len(extract_rows(d))==1 and extract_count(d)==1


def test_xml_recursive():
    assert '<HDon' in find_xml({'detail': {'fileXml':'<HDon><TTChung/></HDon>'}})


def test_key_stable():
    a={'nbmst':'0101234567','khhdon':'AA/26E','shdon':'123','khmshdon':'1'}
    assert invoice_key(a,'sold','hddt') == invoice_key(dict(a),'sold','hddt')


def test_collect_dedup_and_range():
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test')
        from src.invoice.models import Window
        w=Window(date(2026,7,1),date(2026,7,31))
        row={'nbmst':'0101234567','khhdon':'AA','shdon':'1','khmshdon':'1','tdlap':'2026-07-10T10:00:00','tgtcthue':100,'tgtthue':10,'tgtttbso':110}
        storage.upsert_invoice('sold','hddt',w,row,detail={'tdlap':'2026-07-10T10:00:00','tgtcthue':100,'tgtthue':10,'tgtttbso':110},detail_status='OK')
        storage.upsert_invoice('sold','hddt',w,row,detail_status='OK')
        p,s=collect(root,date(2026,7,1),date(2026,7,31))
        assert len(p)==0 and len(s)==1 and s[0]['Tổng cộng']==110


def test_scanner_merges_export_missing_and_detail():
    class Resp:
        def __init__(self,status=200,data=None,content=b''): self.status_code=status; self._data=data; self.content=content; self.headers={}
        def json(self): return self._data
    class Client:
        def __init__(self): self.calls=[]
        def request(self,method,url,**kw):
            self.calls.append((method,url,kw))
            if 'detail' in url:
                return Resp(200,{'tdlap':'2026-07-01T10:00:00','hdhhdvu':[{'mhhdvu':'A'}],'tgtcthue':100,'tgtthue':10,'tgtttbso':110})
            return Resp(200,{'datas':[{'nbmst':'0101234567','khhdon':'AA','shdon':'1','khmshdon':'1','tdlap':'2026-07-01T10:00:00','tgtcthue':100,'tgtthue':10,'tgtttbso':110}], 'totalElements':1})
    with TemporaryDirectory() as td:
        root=Path(td); cfg={'connection':{'base_url':'https://x','invoice_timeout_seconds':2},'invoice_query':{'max_window_days':31,'page_start':0,'max_pages':2,'size':15,'date':{'format':'%d/%m/%YT%H:%M:%S'},'search_template':'tdlap=ge={from_date};tdlap=le={to_date}','sort':{'purchase':'tdlap:desc','sold':'tdlap:desc'},'include_page':True,'throttle_retry_seconds':0},'invoice':{'sold':{'hddt':{'enabled':True,'endpoint':'/api/query/invoices/sold','detail':{'enabled':True,'endpoint':'/api/query/invoices/detail','retries':1,'delay_seconds':0,'backoff_seconds':0},'export':{'enabled':False}}},'purchase':{}}}
        storage=InvoiceStorage(root,'test'); scanner=InvoiceScanner(cfg,Client(),storage,logger=type('L',(),{'info':lambda *a:None,'warning':lambda *a:None,'error':lambda *a:None})())
        result=scanner.run(date(2026,7,1),date(2026,7,1))
        assert result['summary']['sold_hddt']['details_ok']==1
        assert list(root.glob('**/records/*.json'))


def test_money_parsing_and_excel_no_decimal_format():
    from src.reports.excel import money, build_excel
    assert money('1.234.567,89') == 1234568
    assert money('1,234,567.89') == 1234568
    assert money('1.234.567') == 1234567
    assert money('0.00') == 0
    with TemporaryDirectory() as td:
        root = Path(td)
        storage = InvoiceStorage(root, 'test')
        from src.invoice.models import Window
        w = Window(date(2026, 7, 1), date(2026, 7, 31))
        row = {'nbmst':'0101234567','khhdon':'AA','shdon':'1','khmshdon':'1','tdlap':'2026-07-10T10:00:00',
               'tgtcthue':'1.234.567,89','tgtthue':'123.456,78','tgtttbso':'1.358.024,67'}
        storage.upsert_invoice('sold','hddt',w,row,detail_status='OK')
        out = build_excel({'mst':'0101234567','company_name':'TEST'}, root, date(2026,7,1), date(2026,7,31))
        from openpyxl import load_workbook
        wb = load_workbook(out, data_only=False)
        for ws in wb.worksheets:
            for rr in ws.iter_rows():
                for cell in rr:
                    assert cell.number_format != '0.00'
                    assert cell.number_format != '#,##0.00'
                    if isinstance(cell.value, str):
                        assert cell.value.strip() not in {'0.00','0,00'}
        assert wb['Tổng hợp']['D8'].number_format == '#,##0;[Red]-#,##0;0'
        wb.close()


def test_storage_never_downgrades_xml_or_stale_error():
    with TemporaryDirectory() as td:
        root = Path(td); storage = InvoiceStorage(root, 'test')
        from src.invoice.models import Window
        w = Window(date(2026,7,1), date(2026,7,31))
        row = {'nbmst':'0101234567','khhdon':'AA','shdon':'1','khmshdon':'1'}
        storage.upsert_invoice('sold','hddt',w,row,detail={'fileXml':'<HDon><TTChung/></HDon>'},detail_status='OK',error='')
        storage.upsert_invoice('sold','hddt',w,row,detail=None,detail_status='OK',error='')
        files = list(root.glob('**/records/*.json'))
        data = json.loads(files[0].read_text(encoding='utf8'))
        assert data['xml_status'] == 'OK'
        assert data['error'] == ''
        assert data.get('xml_file')


def test_detail_http_200_application_error_is_not_ok():
    class Resp:
        status_code = 200
        headers = {}
        def json(self): return {'message':'Khoảng thời gian không hợp lệ'}
    class Client:
        def request(self,*args,**kwargs): return Resp()
    with TemporaryDirectory() as td:
        root=Path(td)
        cfg={'connection':{'base_url':'https://x','invoice_timeout_seconds':1},'invoice_query':{},
             'invoice':{'sold':{'hddt':{'detail':{'enabled':True,'endpoint':'/api/query/invoices/detail','retries':1,'delay_seconds':0,'backoff_seconds':0}}}}}
        storage=InvoiceStorage(root,'test')
        scanner=InvoiceScanner(cfg,Client(),storage,logger=type('L',(),{})())
        from src.invoice.models import Window
        row={'nbmst':'0101234567','khhdon':'AA','shdon':'1','khmshdon':'1'}
        detail,status,error=scanner.fetch_detail('sold','hddt',cfg['invoice']['sold']['hddt'],Window(date(2026,7,1),date(2026,7,1)),row)
        assert detail is None and status == 'FAILED' and 'application_error' in error


def test_calendar_periods_leap_year_q2_and_q1():
    from src.gui import period_range
    assert period_range(2024,'Q1') == (date(2024,1,1), date(2024,3,31))
    assert period_range(2024,'Q2') == (date(2024,4,1), date(2024,6,30))
    assert (period_range(2024,'Q2')[1] - period_range(2024,'Q2')[0]).days + 1 == 91
    assert period_range(2028,'Q2') == (date(2028,4,1), date(2028,6,30))
    assert period_range(2028,'Cả năm') == (date(2028,1,1), date(2028,12,31))
    assert (period_range(2028,'Cả năm')[1] - period_range(2028,'Cả năm')[0]).days + 1 == 366
    import pytest
    with pytest.raises(ValueError):
        period_range(2026, 'NOPE')

def test_calendar_month_windows_handle_leap_day():
    windows = split_calendar_months(date(2024,1,1), date(2024,6,30))
    assert [(w.start,w.end) for w in windows] == [
        (date(2024,1,1),date(2024,1,31)),
        (date(2024,2,1),date(2024,2,29)),
        (date(2024,3,1),date(2024,3,31)),
        (date(2024,4,1),date(2024,4,30)),
        (date(2024,5,1),date(2024,5,31)),
        (date(2024,6,1),date(2024,6,30)),
    ]

def test_period_ranges_and_full_year_excel_name():
    from src.gui import period_range
    assert period_range(2026,'Q1') == (date(2026,1,1),date(2026,3,31))
    assert period_range(2026,'Q4') == (date(2026,10,1),date(2026,12,31))
    assert period_range(2026,'Cả năm') == (date(2026,1,1),date(2026,12,31))
    assert quarter(date(2026,1,1),date(2026,12,31)) == 'NAM'


def test_collect_filters_current_mst_when_company_changes():
    with TemporaryDirectory() as td:
        root=Path(td); from src.invoice.models import Window
        storage=InvoiceStorage(root,'test',owner_mst='0100000000')
        w=Window(date(2024,4,1),date(2024,4,30))
        mine={'nbmst':'0100000000','nmmst':'0900000000','khhdon':'AA','shdon':'1','khmshdon':'1','tdlap':'2024-04-10T10:00:00','tgtcthue':100,'tgtthue':10,'tgtttbso':110}
        other={'nbmst':'0200000000','nmmst':'0900000000','khhdon':'AA','shdon':'2','khmshdon':'1','tdlap':'2024-04-10T10:00:00','tgtcthue':200,'tgtthue':20,'tgtttbso':220}
        storage.upsert_invoice('sold','hddt',w,mine,detail_status='OK')
        storage.upsert_invoice('sold','hddt',w,other,detail_status='OK')
        p,s=collect(root,date(2024,4,1),date(2024,4,30),owner_mst='0100000000')
        assert len(s)==1 and s[0]['Số hóa đơn']=='1'


def test_storage_does_not_downgrade_successful_detail():
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test',owner_mst='0100000000')
        from src.invoice.models import Window
        w=Window(date(2024,4,1),date(2024,4,30))
        row={'nbmst':'0100000000','khhdon':'AA','shdon':'1','khmshdon':'1'}
        storage.upsert_invoice('sold','hddt',w,row,detail={'ok':True},detail_status='OK',error='')
        fp=storage.upsert_invoice('sold','hddt',w,row,detail=None,detail_status='FAILED',error='temporary')
        data=json.loads(fp.read_text(encoding='utf8'))
        assert data['detail_status']=='OK' and data['detail']=={'ok':True} and data['error']==''

def test_detail_json_is_persisted_and_xml_is_preserved():
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test')
        from src.invoice.models import Window
        w=Window(date(2024,2,1),date(2024,2,29))
        row={'nbmst':'0100000000','khhdon':'AA','shdon':'29','khmshdon':'1'}
        detail={'TTChung': {'SHDon':'29'}, 'fileXml':'<HDon><TTChung><SHDon>29</SHDon></TTChung></HDon>'}
        fp=storage.upsert_invoice('sold','hddt',w,row,detail=detail,detail_status='OK')
        data=json.loads(fp.read_text(encoding='utf8'))
        assert data['detail_file'] and (root/data['detail_file']).exists()
        assert data['xml_file'] and (root/data['xml_file']).read_text(encoding='utf8').startswith('<HDon>')


def test_excel_derives_total_from_prevat_plus_vat_and_flags_source_mismatch():
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test',owner_mst='0100000000')
        from src.invoice.models import Window
        w=Window(date(2024,2,1),date(2024,2,29))
        row={'nbmst':'0100000000','khhdon':'AA','shdon':'29','khmshdon':'1','tdlap':'2024-02-29T10:00:00',
             'tgtcthue':100,'tgtthue':10,'tgtttbso':999}
        storage.upsert_invoice('sold','hddt',w,row,detail_status='OK')
        from src.reports.excel import build_excel
        out=build_excel({'mst':'0100000000','company_name':'TEST'},root,date(2024,2,1),date(2024,2,29))
        from openpyxl import load_workbook
        wb=load_workbook(out,data_only=False)
        ws=wb['Bán']
        # 100 + 10, not the source's deliberately wrong 999.
        assert ws['M4'].value == 110
        assert ws['T4'].value.startswith('NGUỒN=999')
        wb.close()


def test_query_url_encodes_search_expression():
    from src.invoice.scanner import InvoiceScanner
    cfg={'connection':{'base_url':'https://gdt.example'},'invoice_query':{
        'size':15,'page_start':0,'send_page_zero':False,'include_page':True,
        'date':{'format':'%d/%m/%YT%H:%M:%S'},
        'search_template':'tdlap=ge={from_date};tdlap=le={to_date}',
        'sort':{'sold':'tdlap:desc'}
    }}
    scanner=InvoiceScanner(cfg,None,None,None)
    url=scanner._query_url({'endpoint':'/api/query/invoices/sold'},'sold',__import__('src.invoice.models',fromlist=['Window']).Window(date(2026,7,1),date(2026,7,31)),0)
    assert 'search=tdlap%3Dge%3D' in url and '%3B' in url


def test_manifest_completed_after_run():
    class Resp:
        status_code=200; headers={}; content=b''
        def __init__(self,data): self.data=data
        def json(self): return self.data
    class Client:
        def request(self,method,url,**kw):
            return Resp({'datas':[],'totalElements':0})
    with TemporaryDirectory() as td:
        root=Path(td)
        cfg={'connection':{'base_url':'https://x','invoice_timeout_seconds':1},'invoice_query':{
            'max_window_days':31,'page_start':0,'max_pages':2,'size':15,
            'date':{'format':'%d/%m/%YT%H:%M:%S'},'search_template':'tdlap=ge={from_date};tdlap=le={to_date}',
            'sort':{'purchase':'tdlap:desc','sold':'tdlap:desc'},'include_page':True,'throttle_retry_seconds':0},
            'invoice':{'sold':{'hddt':{'enabled':True,'endpoint':'/api/query/invoices/sold','detail':{'enabled':False},'export':{'enabled':False}}},'purchase':{}}}
        storage=InvoiceStorage(root,'test')
        log=type('L',(),{'info':lambda *a:None,'warning':lambda *a:None,'error':lambda *a:None})()
        scanner=InvoiceScanner(cfg,Client(),storage,log)
        scanner.run(date(2026,7,1),date(2026,7,1))
        manifest=json.loads((root/'_diagnostic'/'manifest.json').read_text())
        assert manifest['status']=='COMPLETED'


def test_excel_has_reconciliation_block_and_three_sheets():
    from src.reports.excel import build_excel
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test',owner_mst='0100000000')
        from src.invoice.models import Window
        w=Window(date(2024,2,1),date(2024,2,29))
        storage.upsert_invoice('sold','hddt',w,{'nbmst':'0100000000','khhdon':'AA','shdon':'1','khmshdon':'1','tdlap':'2024-02-10T10:00:00','tgtcthue':100,'tgtthue':10,'tgtttbso':110},detail_status='OK')
        out=build_excel({'mst':'0100000000','company_name':'TEST'},root,date(2024,2,1),date(2024,2,29))
        from openpyxl import load_workbook
        wb=load_workbook(out,data_only=False)
        assert wb.sheetnames == ['Tổng hợp','Mua','Bán']
        assert wb['Tổng hợp']['A13'].value == 'ĐỐI SOÁT DỮ LIỆU'
        assert wb['Tổng hợp']['B14'].value == 0 and wb['Tổng hợp']['C14'].value == 1
        wb.close()


def test_v312_phase_timeouts_are_separate():
    import json
    cfg=json.loads((Path(__file__).resolve().parents[1]/'config'/'gdt_config.json').read_text(encoding='utf-8'))
    assert cfg['connection']['invoice_timeout_seconds'] == 25
    assert cfg['invoice_query']['page_retries'] == 2
    assert cfg['invoice_query']['window_timeout_seconds'] == 90
    assert cfg['invoice_query']['export_timeout_seconds'] == 25
    assert cfg['invoice_query']['detail_phase_timeout_seconds'] == 120
    for side in cfg['invoice'].values():
        for item in side.values():
            assert item['detail']['retries'] == 2
            assert item['detail']['recovery_retries'] == 1
            assert item['export']['retries'] == 1


def test_scanner_progress_callback_is_called():
    from src.invoice.scanner import InvoiceScanner
    events=[]
    cfg={'connection':{'base_url':'https://x','invoice_timeout_seconds':1},'invoice_query':{
        'max_window_days':31,'page_start':0,'max_pages':2,'size':15,
        'date':{'format':'%d/%m/%YT%H:%M:%S'},'search_template':'tdlap=ge={from_date};tdlap=le={to_date}',
        'sort':{'sold':'tdlap:desc'},'include_page':True,'throttle_retry_seconds':0,
        'page_retries':1,'window_timeout_seconds':5},
        'invoice':{'sold':{'hddt':{'enabled':True,'endpoint':'/api/query/invoices/sold','detail':{'enabled':False},'export':{'enabled':False}}},'purchase':{}}}
    class Resp:
        status_code=200; headers={}; content=b''
        def json(self): return {'datas':[],'totalElements':0}
    class Client:
        def request(self,*a,**kw): return Resp()
    with TemporaryDirectory() as td:
        root=Path(td); storage=InvoiceStorage(root,'test')
        log=type('L',(),{'info':lambda *a:None,'warning':lambda *a:None,'error':lambda *a:None})()
        InvoiceScanner(cfg,Client(),storage,log,progress=events.append).run(date(2026,7,1),date(2026,7,1))
    assert any('SOLD / HDDT' in e for e in events)


def test_detail_stops_after_bounded_retries():
    from src.invoice.scanner import InvoiceScanner
    cfg={'connection':{'base_url':'https://x','invoice_timeout_seconds':1},'invoice_query':{},
         'invoice':{}}
    class Client:
        def request(self,*a,**kw): raise RuntimeError('network down')
    scanner=InvoiceScanner(cfg,Client(),None,type('L',(),{'info':lambda *a:None,'error':lambda *a:None})())
    import time as _time
    t=_time.monotonic()
    data,st,err=scanner.fetch_detail('sold','hddt',{'detail':{'enabled':True,'retries':2,'delay_seconds':0,'backoff_seconds':0}},None,
        {'nbmst':'1','khhdon':'AA','shdon':'1','khmshdon':'1'},deadline=t+2)
    assert data is None and st == 'FAILED'
    assert _time.monotonic()-t < 1
